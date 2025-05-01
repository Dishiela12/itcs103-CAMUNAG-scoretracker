from openpyxl import Workbook
from openpyxl import load_workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter
import tkinter as tk
from tkinter import messagebox

wb = Workbook()
ws = wb.active
ws.title = "Student Score Tracker"

# Headers
ws.append(["Name", "Score", "Remarks"])

# Save the workbook
wb.save("student_scores.xlsx")

# Tkinter Form
window = tk.Tk()
window.title("Student Score Tracker")
window.config(bg="#71A5BD")

tk.Label(window, text="Student Name").grid(row=0, column=0, padx=10, pady=5, sticky="w")
tk.Label(window, text="Score").grid(row=1, column=0, padx=10, pady=5, sticky="w")

studentname_entry = tk.Entry(window)
score_entry = tk.Entry(window)

studentname_entry.grid(row=0, column=1, pady=5)
score_entry.grid(row=1, column=1, pady=5)

def validate_inputs():
    studentname = studentname_entry.get()
    score_text = score_entry.get()

    if not studentname or not score_text:
        messagebox.showerror("Error", "All fields are required!")
        return False

    try:
        float(score_text)
    except ValueError:
        messagebox.showerror("Error", "Score must be a number!")
        return False

    return True

def format_excel(ws):
    # Bold headers
    for cell in ws[1]:
        cell.font = Font(bold=True)

    # Bold the "Average" row if exists
    for row in ws.iter_rows(min_row=2):
        if row[0].value == "Average":
            for cell in row:
                cell.font = Font(bold=True)
            break

    # Adjust column widths
    for col in ws.columns:
        max_length = max(len(str(cell.value)) if cell.value else 0 for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = max_length + 2

def save_to_excel():
    if not validate_inputs():
        return

    studentname = studentname_entry.get()
    score = float(score_entry.get())

    wb = load_workbook("student_scores.xlsx")
    ws = wb["Student Score Tracker"]

    for row in ws.iter_rows(min_row=2):
        if row[0].value == "Average":
            ws.delete_rows(row[0].row)
            break

    remarks = "Pass" if score >= 60 else "Fail"
    ws.append([studentname, score, remarks])

    scores = [cell.value for cell in ws["B"] if isinstance(cell.value, (int, float))]
    if scores:
        average = sum(scores) / len(scores)
        ws.append(["Average", average, ""])

    format_excel(ws)
    wb.save("student_scores.xlsx")

    messagebox.showinfo("Success", "Data saved to Excel!")
    studentname_entry.delete(0, tk.END)
    score_entry.delete(0, tk.END)

data_window = None

def show_data():
    global data_window
    
    wb = load_workbook("student_scores.xlsx")
    ws = wb["Student Score Tracker"]

    # Create or reuse the window
    if data_window is None or not data_window.winfo_exists():
        data_window = tk.Toplevel(window)
        data_window.title("Stored Data")
    else:
        # If the window is already open, clear its contents
        for widget in data_window.winfo_children():
            widget.destroy()

    for i, row in enumerate(ws.iter_rows(values_only=True)):
        for j, value in enumerate(row):
            label = tk.Label(data_window, text=value, borderwidth=0, relief="solid", padx=6, pady=3)
            label.grid(row=i, column=j)

# Buttons
tk.Button(window, text="Submit", command=save_to_excel, width=20, bg="#4298BE", fg="white").grid(row=5, column=0, columnspan=2, pady=10)
tk.Button(window, text="View Stored Data", command=show_data, width=20, bg="#1A7FB7", fg="white").grid(row=6, column=0, columnspan=2, pady=10)

window.mainloop()
